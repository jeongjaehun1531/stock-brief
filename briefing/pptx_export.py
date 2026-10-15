"""카드 PNG → 세로형 PPTX (옵션).

모바일에서 PPT 앱으로 열어도 한 장씩 넘겨볼 수 있도록 슬라이드를 9:16 세로로 만들고,
카드 이미지를 폭에 맞춰 배치한다. (내용이 긴 카드는 비율에 맞게 축소)
"""
from __future__ import annotations

from pathlib import Path


def build_pptx(card_paths: list[Path], out_path: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu
    from PIL import Image

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(6858000), Emu(12192000)  # 7.5in × 13.33in (9:16)
    blank = prs.slide_layouts[6]
    sw, sh = prs.slide_width, prs.slide_height
    for p in card_paths:
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(0x0A, 0x0D, 0x14)
        w, h = Image.open(p).size
        ratio = min(sw / w, sh / h)
        pw, ph = int(w * ratio), int(h * ratio)
        slide.shapes.add_picture(str(p), (sw - pw) // 2, (sh - ph) // 2, pw, ph)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path
